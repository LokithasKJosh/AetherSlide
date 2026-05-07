import json
import os
import re
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, AsyncGenerator, Dict, List, Optional

import httpx
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel, Field
import pdfplumber
from PyPDF2 import PdfReader
from pptx import Presentation

try:
    from backend.rag_engine import RagConfig, RagEngine
except ModuleNotFoundError:
    from rag_engine import RagConfig, RagEngine

BASE_DIR = Path(__file__).resolve().parent.parent
UPLOAD_DIR = BASE_DIR / "uploads"
TEMPLATES_DIR = BASE_DIR / "templates"
TEMPLATE_UPLOAD_DIR = TEMPLATES_DIR / "user_uploaded"
GENERATED_DIR = BASE_DIR / "generated"
PUBLIC_DIR = BASE_DIR / "public"
ASSETS_DIR = PUBLIC_DIR / "assets"
TEMPLATE_PREVIEW_DIR = ASSETS_DIR / "template-previews"
GENERATED_PREVIEW_DIR = ASSETS_DIR / "generated-previews"
FRONTEND_DIST = BASE_DIR / "frontend" / "dist"
FRONTEND_WEB_ASSETS = FRONTEND_DIST / "web-assets"

for folder in [UPLOAD_DIR, TEMPLATES_DIR, TEMPLATE_UPLOAD_DIR, GENERATED_DIR, ASSETS_DIR, TEMPLATE_PREVIEW_DIR, GENERATED_PREVIEW_DIR]:
    folder.mkdir(parents=True, exist_ok=True)


def load_local_env(env_path: Path) -> None:
    if not env_path.exists():
        return
    for raw in env_path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and (key not in os.environ or not os.environ.get(key, "").strip()):
            os.environ[key] = value


load_local_env(BASE_DIR / ".env")
os.environ.setdefault("ANONYMIZED_TELEMETRY", "False")

SERVICE_BASE_URL = os.getenv("BASE_URL", "http://127.0.0.1:8001")
LLM_BASE_URL = os.getenv("CUSTOM_LLM_URL") or os.getenv("LLM_URL") or "http://127.0.0.1:8001/v1"
API_KEY = os.getenv("CUSTOM_LLM_API_KEY") or os.getenv("API_KEY", "")
MODEL_NAME = os.getenv("CUSTOM_MODEL") or os.getenv("MODEL_NAME", "qwen32b")
REQUEST_TIMEOUT = float(os.getenv("REQUEST_TIMEOUT", "300"))

DEFAULT_RAG_SOURCE_FILES = [
    Path(r"C:\Users\Lokithas\Desktop\低分子量聚丙烯酸的合成及其阻垢性能研究_张梦豪.pdf"),
    Path(r"C:\Users\Lokithas\Desktop\化工工艺设计手册·下册 (中石化上海工程有限公司) (Z-Library).md"),
]
RAG_ENABLED = os.getenv("RAG_ENABLED", "1").strip() == "1"
RAG_TOP_K = max(1, int(os.getenv("RAG_TOP_K", "3")))
RAG_CHUNK_SIZE = max(256, int(os.getenv("RAG_CHUNK_SIZE", "1000")))
RAG_CHUNK_OVERLAP = max(0, int(os.getenv("RAG_CHUNK_OVERLAP", "180")))
RAG_COLLECTION_NAME = os.getenv("RAG_COLLECTION_NAME", "ppt_rag_knowledge")
RAG_STORE_DIR = BASE_DIR / "rag_store"

raw_rag_sources = os.getenv("RAG_SOURCE_FILES", "").strip()
if raw_rag_sources:
    parsed_sources = [s.strip() for s in re.split(r"[|;\n]", raw_rag_sources) if s.strip()]
    RAG_SOURCE_FILES = [Path(item) for item in parsed_sources]
else:
    RAG_SOURCE_FILES = DEFAULT_RAG_SOURCE_FILES

ALLOWED_DOC_EXTENSIONS = {".txt", ".docx", ".pdf"}
ALLOWED_TEMPLATE_EXTENSIONS = {".pptx", ".potx"}
UPLOAD_CONTEXT: Dict[str, str] = {}

SYSTEM_PROMPT = """你是一个擅长将文档内容转成 PPT 大纲的助手。
请先给出简短的生成过程说明，然后务必输出一段 JSON（用于生成 PPT）。
JSON 只使用这个结构：
{
  "title": "整体标题",
  "slides": [
    {
      "title": "页面标题",
      "bullets": ["要点1", "要点2", "要点3"]
    }
  ]
}
请保证 JSON 可被标准 json.loads 解析。
"""


class ChatRequest(BaseModel):
    user_input: str = Field(..., min_length=1)
    file_id: Optional[str] = None
    document_text: Optional[str] = None
    template_name: Optional[str] = None


app = FastAPI(title="AI PPT MVP", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

RAG_ENGINE: Optional[RagEngine] = None
RAG_BOOT_ERROR = ""
if RAG_ENABLED:
    try:
        RAG_ENGINE = RagEngine(
            RagConfig(
                persist_dir=RAG_STORE_DIR,
                collection_name=RAG_COLLECTION_NAME,
                source_files=RAG_SOURCE_FILES,
                chunk_size=RAG_CHUNK_SIZE,
                chunk_overlap=RAG_CHUNK_OVERLAP,
            )
        )
    except Exception as rag_exc:  # noqa: BLE001
        RAG_BOOT_ERROR = str(rag_exc)


def build_rag_context(query_text: str, top_k: int) -> Dict[str, Any]:
    if not RAG_ENABLED:
        return {"enabled": False, "items": [], "error": ""}
    if RAG_BOOT_ERROR:
        return {"enabled": True, "items": [], "error": RAG_BOOT_ERROR}
    if RAG_ENGINE is None:
        return {"enabled": True, "items": [], "error": "RAG engine unavailable"}

    try:
        hits = RAG_ENGINE.retrieve(query_text, top_k=top_k)
        return {"enabled": True, "items": hits, "error": ""}
    except Exception as rag_exc:  # noqa: BLE001
        return {"enabled": True, "items": [], "error": str(rag_exc)}


def sse_event(event: str, data: Dict[str, Any]) -> str:
    payload = json.dumps(data, ensure_ascii=False)
    return f"event: {event}\ndata: {payload}\n\n"


def resolve_chat_completions_url(base_url: str) -> str:
    normalized = (base_url or "").strip().rstrip("/")
    if not normalized:
        raise RuntimeError("CUSTOM_LLM_URL is empty. Please configure environment variables.")
    if normalized.endswith("/chat/completions"):
        return normalized
    if normalized.endswith("/v1"):
        return f"{normalized}/chat/completions"
    return f"{normalized}/v1/chat/completions"


def sanitize_filename(raw_name: str, max_len: int = 48) -> str:
    name = (raw_name or "").strip()
    name = re.sub(r"\s+", "_", name)
    name = re.sub(r'[\\/:*?"<>|]+', "", name)
    name = re.sub(r"[^\w\u4e00-\u9fff-]+", "", name, flags=re.UNICODE)
    name = name[:max_len].strip("._ ")
    return name or "template"


def build_unique_path(directory: Path, stem: str, suffix: str) -> Path:
    candidate = directory / f"{stem}{suffix}"
    if not candidate.exists():
        return candidate
    index = 2
    while True:
        candidate = directory / f"{stem}_{index}{suffix}"
        if not candidate.exists():
            return candidate
        index += 1


def get_font(size: int) -> ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/simsun.ttc",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def draw_preview_card(title: str, subtitle: str, lines: List[str], output_path: Path) -> None:
    width, height = 1280, 720
    image = Image.new("RGB", (width, height), (22, 20, 34))
    draw = ImageDraw.Draw(image)
    title_font = get_font(56)
    subtitle_font = get_font(28)
    body_font = get_font(34)

    draw.rectangle((40, 40, width - 40, height - 40), outline=(180, 170, 230), width=3)
    draw.text((70, 80), title, fill=(240, 240, 250), font=title_font)
    draw.text((70, 160), subtitle, fill=(190, 185, 220), font=subtitle_font)

    y = 240
    for line in lines[:8]:
        draw.text((90, y), f"- {line}", fill=(230, 230, 245), font=body_font)
        y += 56

    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, format="PNG")


def extract_first_slide_lines(ppt_file: Path, max_lines: int = 8) -> List[str]:
    lines: List[str] = []
    try:
        prs = Presentation(str(ppt_file))
    except Exception:  # noqa: BLE001
        return lines
    if not prs.slides:
        return lines

    first_slide = prs.slides[0]
    for shape in first_slide.shapes:
        text = getattr(shape, "text", "")
        if not text:
            continue
        for line in text.splitlines():
            clean = line.strip()
            if clean:
                lines.append(clean)
            if len(lines) >= max_lines:
                return lines
    return lines


def extract_all_slide_payloads(ppt_file: Path, max_lines: int = 10) -> List[Dict[str, Any]]:
    payloads: List[Dict[str, Any]] = []
    try:
        prs = Presentation(str(ppt_file))
    except Exception:  # noqa: BLE001
        return payloads

    for index, slide in enumerate(prs.slides, start=1):
        title_text = ""
        if slide.shapes.title and getattr(slide.shapes.title, "text", ""):
            title_text = slide.shapes.title.text.strip()

        lines: List[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if not text:
                continue
            for raw in text.splitlines():
                clean = re.sub(r"^[\s\-•·]+", "", raw).strip()
                if clean:
                    lines.append(clean)
                if len(lines) >= max_lines:
                    break
            if len(lines) >= max_lines:
                break

        if title_text and lines and lines[0] == title_text:
            lines = lines[1:]
        if not lines:
            lines = ["No extractable text on this slide."]
        if not title_text:
            title_text = f"Slide {index}"

        payloads.append({"title": title_text, "lines": lines})
    return payloads


def ensure_generated_ppt_previews(ppt_file: Path) -> List[str]:
    slide_payloads = extract_all_slide_payloads(ppt_file)
    if not slide_payloads:
        slide_payloads = [{"title": ppt_file.stem, "lines": ["No preview content available."]}]

    preview_folder_name = sanitize_filename(ppt_file.stem, 64)
    preview_folder = GENERATED_PREVIEW_DIR / preview_folder_name
    preview_folder.mkdir(parents=True, exist_ok=True)

    source_mtime = int(ppt_file.stat().st_mtime)
    manifest_path = preview_folder / "manifest.json"
    valid_cache = False

    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            slide_count = int(manifest.get("slide_count", 0))
            manifest_mtime = int(manifest.get("source_mtime", 0))
            if manifest_mtime == source_mtime and slide_count == len(slide_payloads):
                valid_cache = True
                for i in range(1, slide_count + 1):
                    if not (preview_folder / f"slide_{i:03d}.png").exists():
                        valid_cache = False
                        break
        except Exception:  # noqa: BLE001
            valid_cache = False

    if not valid_cache:
        for old_png in preview_folder.glob("slide_*.png"):
            try:
                old_png.unlink()
            except OSError:
                pass

        total = len(slide_payloads)
        for index, payload in enumerate(slide_payloads, start=1):
            output_path = preview_folder / f"slide_{index:03d}.png"
            draw_preview_card(
                title=str(payload.get("title", f"Slide {index}")),
                subtitle=f"Slide {index}/{total}",
                lines=[str(line) for line in payload.get("lines", [])],
                output_path=output_path,
            )

        manifest_path.write_text(
            json.dumps({"source_mtime": source_mtime, "slide_count": len(slide_payloads)}, ensure_ascii=False),
            encoding="utf-8",
        )

    return [
        f"/assets/generated-previews/{preview_folder_name}/slide_{idx:03d}.png?t={source_mtime}"
        for idx in range(1, len(slide_payloads) + 1)
    ]


def ensure_pptx_preview(ppt_file: Path, namespace: str) -> str:
    preview_name = sanitize_filename(f"{namespace}_{ppt_file.stem}", 80) + ".png"
    preview_path = TEMPLATE_PREVIEW_DIR / preview_name

    needs_regen = not preview_path.exists() or preview_path.stat().st_mtime < ppt_file.stat().st_mtime
    if needs_regen:
        lines = extract_first_slide_lines(ppt_file)
        if not lines:
            lines = ["第一页未解析到文本内容。", "请上传包含文本内容的 PPT 模板。"]
        draw_preview_card(
            title=ppt_file.stem,
            subtitle="第一页预览",
            lines=lines,
            output_path=preview_path,
        )
    return f"/assets/template-previews/{preview_name}"


def ensure_bundle_preview(bundle_dir: Path) -> str:
    preview_name = sanitize_filename(f"local_{bundle_dir.name}", 80) + ".png"
    preview_path = TEMPLATE_PREVIEW_DIR / preview_name
    if preview_path.exists():
        return f"/assets/template-previews/{preview_name}"

    settings_path = bundle_dir / "settings.json"
    description = ""
    if settings_path.exists():
        try:
            settings = json.loads(settings_path.read_text(encoding="utf-8", errors="ignore"))
            description = str(settings.get("description", "")).strip()
        except json.JSONDecodeError:
            description = ""

    first_layout = ""
    tsx_files = sorted(bundle_dir.glob("*.tsx"))
    if tsx_files:
        first_layout = tsx_files[0].stem

    lines = []
    if description:
        lines.append(description)
    if first_layout:
        lines.append(f"首个布局：{first_layout}")
    if not lines:
        lines = ["本地模板包", "该模板为布局集合。"]

    draw_preview_card(
        title=bundle_dir.name,
        subtitle="本地模板",
        lines=lines,
        output_path=preview_path,
    )
    return f"/assets/template-previews/{preview_name}"


def extract_text_from_file(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext == ".txt":
        return file_path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".docx":
        doc = Document(str(file_path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if ext == ".pdf":
        pages: List[str] = []
        try:
            reader = PdfReader(str(file_path))
            pages = [(page.extract_text() or "").strip() for page in reader.pages]
        except Exception:  # noqa: BLE001
            pages = []

        extracted = "\n".join(page for page in pages if page).strip()
        if extracted:
            return extracted

        try:
            with pdfplumber.open(str(file_path)) as pdf:
                alt_pages = [(page.extract_text() or "").strip() for page in pdf.pages]
            return "\n".join(page for page in alt_pages if page)
        except Exception:  # noqa: BLE001
            return ""
    raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")


def extract_json_from_text(text: str) -> Optional[Dict[str, Any]]:
    fenced = re.search(r"```json\s*(\{[\s\S]*?\})\s*```", text, re.DOTALL)
    if fenced:
        try:
            parsed = json.loads(fenced.group(1).strip())
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            pass

    decoder = json.JSONDecoder()
    for idx, ch in enumerate(text):
        if ch != "{":
            continue
        try:
            parsed, _ = decoder.raw_decode(text[idx:])
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            return parsed
    return None


def fallback_ppt_json(user_input: str, document_text: str) -> Dict[str, Any]:
    lines = [line.strip("-• ").strip() for line in document_text.splitlines() if line.strip()]
    key_points = lines[:6] if lines else ["未提取到有效文档内容，使用用户要求生成初稿。"]
    return {
        "title": "AI 生成演示文稿",
        "slides": [
            {"title": "需求概述", "bullets": [user_input]},
            {"title": "参考文档要点", "bullets": key_points},
        ],
    }


def resolve_template_for_generation(template_name: Optional[str]) -> Optional[Path]:
    if not template_name:
        return None
    candidate = (TEMPLATES_DIR / template_name).resolve()
    try:
        candidate.relative_to(TEMPLATES_DIR.resolve())
    except ValueError:
        return None
    if candidate.exists() and candidate.is_file() and candidate.suffix.lower() in ALLOWED_TEMPLATE_EXTENSIONS:
        return candidate
    return None


def generate_ppt(
    ppt_data: Dict[str, Any],
    template_name: Optional[str] = None,
    file_name_hint: Optional[str] = None,
) -> Path:
    template_path = resolve_template_for_generation(template_name)
    presentation = Presentation(str(template_path)) if template_path else Presentation()

    slides = ppt_data.get("slides") or []
    if not slides:
        slides = [{"title": ppt_data.get("title", "Untitled"), "bullets": ["No content generated."]}]

    for item in slides:
        layout = presentation.slide_layouts[1] if len(presentation.slide_layouts) > 1 else presentation.slide_layouts[0]
        slide = presentation.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = str(item.get("title", "Untitled"))

        bullets = item.get("bullets") or []
        bullet_lines = [str(point).strip() for point in bullets if str(point).strip()]
        body_text = "\n".join(f"- {line}" for line in bullet_lines) or " "
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = body_text

    stem = sanitize_filename(file_name_hint or "generated_ppt", 48)
    output_path = build_unique_path(GENERATED_DIR, stem, ".pptx")
    presentation.save(str(output_path))
    return output_path


async def stream_qwen(messages: List[Dict[str, str]]) -> AsyncGenerator[str, None]:
    if not API_KEY:
        raise RuntimeError("CUSTOM_LLM_API_KEY is empty. Please configure environment variables.")

    url = resolve_chat_completions_url(LLM_BASE_URL)
    headers = {"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": MODEL_NAME,
        "messages": messages,
        "temperature": 0.4,
        "stream": True,
    }

    timeout = httpx.Timeout(REQUEST_TIMEOUT, connect=30.0)
    async with httpx.AsyncClient(timeout=timeout) as client:
        async with client.stream("POST", url, headers=headers, json=payload) as response:
            response.raise_for_status()
            async for line in response.aiter_lines():
                if not line or not line.startswith("data:"):
                    continue
                data = line[5:].strip()
                if data == "[DONE]":
                    break
                try:
                    chunk = json.loads(data)
                except json.JSONDecodeError:
                    continue
                token = chunk.get("choices", [{}])[0].get("delta", {}).get("content")
                if token:
                    yield token


@app.on_event("startup")
async def startup_prepare_rag() -> None:
    if not RAG_ENABLED or RAG_BOOT_ERROR or RAG_ENGINE is None:
        return
    try:
        RAG_ENGINE.ensure_index_built(force=False)
    except Exception as rag_exc:  # noqa: BLE001
        print(f"[RAG] startup build failed: {rag_exc}")


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)) -> Dict[str, Any]:
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_DOC_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Only {sorted(ALLOWED_DOC_EXTENSIONS)} are supported.")

    stored_path = UPLOAD_DIR / f"{uuid.uuid4().hex}{ext}"
    with stored_path.open("wb") as out:
        shutil.copyfileobj(file.file, out)

    text = extract_text_from_file(stored_path)
    file_id = uuid.uuid4().hex
    UPLOAD_CONTEXT[file_id] = text

    return {
        "file_id": file_id,
        "filename": filename,
        "char_count": len(text),
        "preview": text[:1000],
    }


@app.post("/api/chat")
async def chat(payload: ChatRequest) -> StreamingResponse:
    async def event_generator() -> AsyncGenerator[str, None]:
        document_text = payload.document_text or UPLOAD_CONTEXT.get(payload.file_id or "", "")
        rag_query_seed = f"{payload.user_input}\n{document_text[:2000]}".strip()
        rag_result = build_rag_context(rag_query_seed, top_k=RAG_TOP_K)
        rag_hits = rag_result.get("items", [])
        rag_error = rag_result.get("error", "")
        rag_context_text = "\n\n".join(
            [
                f"[{idx + 1}] Source: {item.get('source_name') or 'unknown'} (chunk {item.get('chunk_index', 0)})\n{item.get('text', '')}"
                for idx, item in enumerate(rag_hits)
            ]
        )
        if not rag_context_text:
            rag_context_text = "No retrieved context."
        combined_prompt = (
            f"用户要求：\n{payload.user_input}\n\n"
            f"RAG 检索结果（Top {RAG_TOP_K}，如相关请优先参考）：\n{rag_context_text}\n\n"
            f"用户上传文档文本（可能为空）：\n{document_text[:20000]}"
        )
        messages = [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": combined_prompt},
        ]

        full_text = ""
        try:
            async for token in stream_qwen(messages):
                full_text += token
                yield sse_event("token", {"text": token})

            parsed = extract_json_from_text(full_text)
            ppt_data = parsed or fallback_ppt_json(payload.user_input, document_text)
            output_file = generate_ppt(
                ppt_data=ppt_data,
                template_name=payload.template_name,
                file_name_hint=payload.user_input,
            )
            preview_images = ensure_generated_ppt_previews(output_file)

            yield sse_event(
                "done",
                {
                    "message": "PPT generated",
                    "ppt_file": output_file.name,
                    "download_url": f"/generated/{output_file.name}",
                    "preview_images": preview_images,
                    "used_fallback_json": parsed is None,
                    "rag": {
                        "enabled": bool(rag_result.get("enabled")),
                        "error": rag_error,
                        "hits": [
                            {
                                "source_name": item.get("source_name", ""),
                                "chunk_index": item.get("chunk_index", 0),
                                "score": round(float(item.get("score", 0.0)), 4),
                            }
                            for item in rag_hits
                        ],
                    },
                },
            )
        except Exception as exc:  # noqa: BLE001
            yield sse_event("error", {"message": str(exc)})

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@app.get("/api/templates")
async def list_templates() -> Dict[str, Any]:
    items: List[Dict[str, Any]] = []

    for folder in sorted([d for d in TEMPLATES_DIR.iterdir() if d.is_dir() and d.name != TEMPLATE_UPLOAD_DIR.name]):
        if not (folder / "settings.json").exists():
            continue
        items.append(
            {
                "name": folder.name,
                "template_file": folder.name,
                "download_url": "",
                "thumbnail_url": "",
                "template_type": "本地模板",
                "preview_unsupported": True,
                "preview_message": "格式不支持预览",
            }
        )

    for ppt_file in sorted(TEMPLATES_DIR.glob("*")):
        if not ppt_file.is_file() or ppt_file.suffix.lower() not in ALLOWED_TEMPLATE_EXTENSIONS:
            continue
        rel = ppt_file.relative_to(TEMPLATES_DIR).as_posix()
        items.append(
            {
                "name": ppt_file.stem,
                "template_file": rel,
                "download_url": f"/template-files/{rel}",
                "thumbnail_url": ensure_pptx_preview(ppt_file, "local"),
                "template_type": "本地模板",
            }
        )

    user_files = sorted(
        [f for f in TEMPLATE_UPLOAD_DIR.glob("*") if f.is_file() and f.suffix.lower() in ALLOWED_TEMPLATE_EXTENSIONS],
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )
    for file_path in user_files:
        rel = file_path.relative_to(TEMPLATES_DIR).as_posix()
        items.append(
            {
                "name": file_path.stem,
                "template_file": rel,
                "download_url": f"/template-files/{rel}",
                "thumbnail_url": ensure_pptx_preview(file_path, "user"),
                "template_type": "用户上传",
            }
        )

    return {"items": items}


@app.post("/api/templates/upload")
async def upload_template(file: UploadFile = File(...)) -> Dict[str, Any]:
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_TEMPLATE_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Only {sorted(ALLOWED_TEMPLATE_EXTENSIONS)} are supported.")

    stem = sanitize_filename(Path(filename).stem, 48)
    target = build_unique_path(TEMPLATE_UPLOAD_DIR, stem, ext)
    with target.open("wb") as out:
        shutil.copyfileobj(file.file, out)

    rel = target.relative_to(TEMPLATES_DIR).as_posix()
    thumb = ensure_pptx_preview(target, "user")
    return {
        "message": "模板上传成功",
        "item": {
            "name": target.stem,
            "template_file": rel,
            "download_url": f"/template-files/{rel}",
            "thumbnail_url": thumb,
            "template_type": "用户上传",
        },
    }


@app.get("/api/history")
async def list_history() -> Dict[str, Any]:
    files = sorted(GENERATED_DIR.glob("*.pptx"), key=lambda item: item.stat().st_mtime, reverse=True)
    items = []
    for item in files:
        stat = item.stat()
        items.append(
            {
                "file_name": item.name,
                "size_bytes": stat.st_size,
                "updated_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                "download_url": f"/generated/{item.name}",
            }
        )
    return {"items": items}


@app.get("/api/rag/status")
async def rag_status() -> Dict[str, Any]:
    if not RAG_ENABLED:
        return {"enabled": False, "error": "", "source_files": [str(p) for p in RAG_SOURCE_FILES], "chunk_count": 0}
    if RAG_BOOT_ERROR:
        return {"enabled": True, "error": RAG_BOOT_ERROR, "source_files": [str(p) for p in RAG_SOURCE_FILES], "chunk_count": 0}
    if RAG_ENGINE is None:
        return {"enabled": True, "error": "RAG engine unavailable", "source_files": [str(p) for p in RAG_SOURCE_FILES], "chunk_count": 0}
    return {
        "enabled": True,
        "error": "",
        "source_files": [str(p) for p in RAG_SOURCE_FILES],
        "chunk_count": int(RAG_ENGINE.collection.count()),
        "top_k": RAG_TOP_K,
        "chunk_size": RAG_CHUNK_SIZE,
        "chunk_overlap": RAG_CHUNK_OVERLAP,
    }


@app.post("/api/rag/rebuild")
async def rag_rebuild() -> Dict[str, Any]:
    if not RAG_ENABLED:
        return {"enabled": False, "rebuilt": False, "message": "RAG is disabled"}
    if RAG_BOOT_ERROR:
        raise HTTPException(status_code=500, detail=RAG_BOOT_ERROR)
    if RAG_ENGINE is None:
        raise HTTPException(status_code=500, detail="RAG engine unavailable")
    result = RAG_ENGINE.ensure_index_built(force=True)
    return {
        "enabled": True,
        "rebuilt": bool(result.get("rebuilt", False)),
        "chunk_count": int(result.get("chunk_count", 0)),
        "source_count": int(result.get("source_count", 0)),
    }


@app.get("/api/generated-preview")
async def generated_preview(ppt_file: str) -> Dict[str, Any]:
    if not ppt_file:
        raise HTTPException(status_code=400, detail="ppt_file is required")

    candidate = (GENERATED_DIR / ppt_file).resolve()
    try:
        candidate.relative_to(GENERATED_DIR.resolve())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid ppt_file path") from exc

    if not candidate.exists() or not candidate.is_file() or candidate.suffix.lower() != ".pptx":
        raise HTTPException(status_code=404, detail="PPT file not found")

    preview_images = ensure_generated_ppt_previews(candidate)
    return {
        "ppt_file": candidate.name,
        "slide_count": len(preview_images),
        "items": [{"index": i + 1, "image_url": url} for i, url in enumerate(preview_images)],
    }


@app.get("/api/health")
async def health() -> Dict[str, str]:
    return {"status": "ok"}


app.mount("/assets", StaticFiles(directory=ASSETS_DIR), name="assets")
app.mount("/generated", StaticFiles(directory=GENERATED_DIR), name="generated")
app.mount("/template-files", StaticFiles(directory=TEMPLATES_DIR), name="template_files")

if FRONTEND_WEB_ASSETS.exists():
    app.mount("/web-assets", StaticFiles(directory=FRONTEND_WEB_ASSETS), name="web_assets")


@app.get("/", include_in_schema=False)
async def root():
    index_file = FRONTEND_DIST / "index.html"
    if index_file.exists():
        return FileResponse(index_file)
    return JSONResponse({"message": "Frontend dist not found. Build frontend first."}, status_code=404)


@app.get("/{full_path:path}", include_in_schema=False)
async def spa_fallback(full_path: str):
    if full_path.startswith("api/"):
        return JSONResponse({"detail": "Not Found"}, status_code=404)
    index_file = FRONTEND_DIST / "index.html"
    if index_file.exists():
        return FileResponse(index_file)
    return JSONResponse({"message": "Frontend dist not found. Build frontend first."}, status_code=404)
